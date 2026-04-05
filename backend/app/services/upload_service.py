import os
import uuid
from pathlib import Path
from typing import Optional

UPLOAD_DIR = os.getenv("UPLOAD_DIR", "./uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


async def save_upload(file) -> dict:
    file_id = str(uuid.uuid4())
    safe_name = Path(file.filename).name.replace(" ", "_")
    dest = os.path.join(UPLOAD_DIR, f"{file_id}_{safe_name}")
    content = await file.read()
    with open(dest, "wb") as f:
        f.write(content)
    return {"file_id": file_id, "filename": file.filename, "path": dest}


def get_upload_path(file_id: str) -> Optional[str]:
    for f in Path(UPLOAD_DIR).glob(f"{file_id}_*"):
        return str(f)
    return None


def extract_pptx_text(path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
        text = "\n".join(parts)
        if len(text) > 3000:
            text = text[:3000] + "\n...[truncated]"
        return text
    except Exception as e:
        return ""
